"""Deterministic file parsing for knowledge ingestion."""

from __future__ import annotations

from pathlib import Path
from typing import Any, Literal, TypedDict
from urllib.parse import quote


BlockType = Literal["text", "table"]
_TEXT_SUFFIXES = {".txt", ".md", ".csv", ".json", ".html", ".xml", ".rtf"}
_IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp"}


class DocumentParseError(ValueError):
    """Raised when a supported document cannot be parsed safely."""


class UnsupportedDocumentTypeError(DocumentParseError):
    """Raised instead of interpreting an unknown binary as text."""


class ParsedBlock(TypedDict, total=False):
    type: BlockType
    text: str
    source_locator: str
    page: int
    slide: int
    sheet: str
    row_count: int
    row_start: int
    row_end: int
    line_start: int
    line_end: int
    metadata: dict[str, Any]


class ParsedDocument(TypedDict):
    filename: str
    mime_type: str
    blocks: list[ParsedBlock]
    metadata: dict[str, Any]


def _read_text(path: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gb18030", "latin-1"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_bytes().decode("utf-8", errors="replace")


def _parse_text(path: Path, lines_per_block: int = 200) -> tuple[list[ParsedBlock], dict[str, Any]]:
    lines = _read_text(path).splitlines()
    blocks: list[ParsedBlock] = []
    for offset in range(0, len(lines), lines_per_block):
        end = min(len(lines), offset + lines_per_block)
        text = "\n".join(lines[offset:end]).strip()
        if not text:
            continue
        start_line = offset + 1
        blocks.append(
            {
                "type": "text",
                "text": text,
                "source_locator": f"text:lines={start_line}-{end}",
                "line_start": start_line,
                "line_end": end,
                "metadata": {"line_start": start_line, "line_end": end},
            }
        )
    return blocks, {"line_count": len(lines)}


def _parse_pdf(path: Path) -> tuple[list[ParsedBlock], dict[str, Any]]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            unlocked = reader.decrypt("")
        except Exception as exc:  # pragma: no cover - parser-specific exception types vary
            raise DocumentParseError("Encrypted PDF requires a password.") from exc
        if not unlocked:
            raise DocumentParseError("Encrypted PDF requires a password.")

    blocks: list[ParsedBlock] = []
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            blocks.append(
                {
                    "type": "text",
                    "text": text,
                    "source_locator": f"pdf:page={index}",
                    "page": index,
                    "metadata": {"page": index},
                }
            )
    return blocks, {"page_count": len(reader.pages)}


def _parse_docx(path: Path) -> tuple[list[ParsedBlock], dict[str, Any]]:
    from docx import Document

    document = Document(str(path))
    blocks: list[ParsedBlock] = []
    for paragraph_index, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = paragraph.style.name if paragraph.style is not None else ""
        blocks.append(
            {
                "type": "text",
                "text": text,
                "source_locator": f"docx:paragraph={paragraph_index}",
                "metadata": {
                    "paragraph_index": paragraph_index,
                    "style": style_name,
                },
            }
        )

    for table_index, table in enumerate(document.tables, start=1):
        rows: list[str] = []
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            blocks.append(
                {
                    "type": "table",
                    "text": "\n".join(rows),
                    "source_locator": f"docx:table={table_index}",
                    "row_count": len(rows),
                    "metadata": {"table_index": table_index},
                }
            )
    return blocks, {
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
    }


def _shape_position(shape: Any, fallback_index: int) -> tuple[int, int, int]:
    return (
        int(getattr(shape, "top", 0) or 0),
        int(getattr(shape, "left", 0) or 0),
        fallback_index,
    )


def _parse_pptx(path: Path) -> tuple[list[ParsedBlock], dict[str, Any]]:
    from pptx import Presentation

    presentation = Presentation(str(path))
    blocks: list[ParsedBlock] = []
    visual_slide_count = 0
    for slide_index, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []
        has_visual = False
        positioned_shapes = sorted(
            enumerate(slide.shapes),
            key=lambda item: _shape_position(item[1], item[0]),
        )
        for shape_index, shape in positioned_shapes:
            if getattr(shape, "has_text_frame", False):
                text = str(getattr(shape, "text", "") or "").strip()
                if text:
                    parts.append(text)
            if getattr(shape, "has_table", False):
                table_rows: list[str] = []
                for row in shape.table.rows:
                    cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    if any(cells):
                        table_rows.append(" | ".join(cells))
                if table_rows:
                    parts.append("\n".join(table_rows))
            if getattr(shape, "shape_type", None) in {3, 13} or getattr(shape, "has_chart", False):
                has_visual = True
        if has_visual:
            visual_slide_count += 1
        text = "\n\n".join(parts).strip()
        if text:
            blocks.append(
                {
                    "type": "text",
                    "text": text,
                    "source_locator": f"pptx:slide={slide_index}",
                    "slide": slide_index,
                    "metadata": {
                        "slide": slide_index,
                        "has_visual": has_visual,
                    },
                }
            )
    return blocks, {
        "slide_count": len(presentation.slides),
        "visual_slide_count": visual_slide_count,
    }


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("|", "\\|").replace("\n", " ").strip()


def _rows_to_markdown(rows: list[list[Any]]) -> str:
    normalized = [[_cell_text(value) for value in row] for row in rows]
    width = max((len(row) for row in normalized), default=0)
    if width == 0:
        return ""
    padded = [row + [""] * (width - len(row)) for row in normalized]
    header = padded[0]
    body = padded[1:]
    lines = ["| " + " | ".join(header) + " |", "| " + " | ".join(["---"] * width) + " |"]
    lines.extend("| " + " | ".join(row) + " |" for row in body)
    return "\n".join(lines)


def _parse_xlsx(path: Path, rows_per_block: int = 200) -> tuple[list[ParsedBlock], dict[str, Any]]:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    blocks: list[ParsedBlock] = []
    try:
        for sheet in workbook.worksheets:
            pending: list[tuple[int, list[Any]]] = []

            def flush() -> None:
                if not pending:
                    return
                start_row = pending[0][0]
                end_row = pending[-1][0]
                rows = [row for _, row in pending]
                max_column = max((len(row) for row in rows), default=1)
                range_ref = f"A{start_row}:{get_column_letter(max_column)}{end_row}"
                encoded_sheet = quote(str(sheet.title), safe="")
                blocks.append(
                    {
                        "type": "table",
                        "text": _rows_to_markdown(rows),
                        "source_locator": f"xlsx:sheet={encoded_sheet}&range={range_ref}",
                        "sheet": str(sheet.title),
                        "row_count": len(rows),
                        "row_start": start_row,
                        "row_end": end_row,
                        "metadata": {
                            "sheet": str(sheet.title),
                            "range": range_ref,
                        },
                    }
                )
                pending.clear()

            for row_index, values in enumerate(sheet.iter_rows(values_only=True), start=1):
                row = list(values)
                while row and row[-1] is None:
                    row.pop()
                if not row or not any(value not in (None, "") for value in row):
                    continue
                pending.append((row_index, row))
                if len(pending) >= rows_per_block:
                    flush()
            flush()
    finally:
        workbook.close()
    return blocks, {"sheet_count": len(workbook.sheetnames)}


def parse_file(path: str | Path, mime_type: str = "", filename: str | None = None) -> ParsedDocument:
    """Parse a supported local file into locator-stable text/table blocks."""

    file_path = Path(path)
    suffix = (Path(filename).suffix if filename else file_path.suffix).lower()
    display_name = filename or file_path.name
    parse_metadata: dict[str, Any]
    route_mode = "full-context"

    if suffix == ".pdf" or mime_type == "application/pdf":
        blocks, parse_metadata = _parse_pdf(file_path)
    elif suffix == ".docx" or mime_type.endswith("wordprocessingml.document"):
        blocks, parse_metadata = _parse_docx(file_path)
    elif suffix == ".pptx" or mime_type.endswith("presentationml.presentation"):
        blocks, parse_metadata = _parse_pptx(file_path)
    elif suffix == ".xlsx" or mime_type.endswith("spreadsheetml.sheet"):
        blocks, parse_metadata = _parse_xlsx(file_path)
    elif suffix in _IMAGE_SUFFIXES or mime_type.startswith("image/"):
        blocks, parse_metadata = [], {}
        route_mode = "vision"
    elif suffix in _TEXT_SUFFIXES or mime_type.startswith("text/") or mime_type in {
        "application/json",
        "application/xml",
        "application/rtf",
    }:
        blocks, parse_metadata = _parse_text(file_path)
        route_mode = "text"
    else:
        raise UnsupportedDocumentTypeError(f"Unsupported document type: {suffix or mime_type or 'unknown'}")

    processing_status = "parsed"
    if not blocks and (route_mode == "vision" or suffix in {".pdf", ".pptx"}):
        processing_status = "vision_required"
        route_mode = "vision"

    return {
        "filename": display_name,
        "mime_type": mime_type,
        "blocks": blocks,
        "metadata": {
            "suffix": suffix,
            "block_count": len(blocks),
            "size_bytes": file_path.stat().st_size,
            "parser": "builtin",
            "route_mode": route_mode,
            "processing_status": processing_status,
            **parse_metadata,
        },
    }
