from pathlib import Path

from docx import Document
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pypdf import PdfWriter

from app.services.file_parser_service import parse_file


def test_pptx_parser_emits_slide_locators(tmp_path: Path):
    path = tmp_path / "deck.pptx"
    presentation = Presentation()
    for title in ("第一页标题", "第二页标题"):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        shape = slide.shapes.add_textbox(100, 100, 500, 100)
        shape.text = title
    presentation.save(path)

    parsed = parse_file(path, filename="deck.pptx")

    assert [block["source_locator"] for block in parsed["blocks"]] == [
        "pptx:slide=1",
        "pptx:slide=2",
    ]
    assert "第一页标题" in parsed["blocks"][0]["text"]
    assert "第二页标题" in parsed["blocks"][1]["text"]
    assert parsed["metadata"]["slide_count"] == 2


def test_docx_parser_emits_paragraph_and_table_locators(tmp_path: Path):
    path = tmp_path / "brief.docx"
    document = Document()
    document.add_paragraph("项目背景")
    document.add_paragraph("")
    document.add_paragraph("交付计划")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "事项"
    table.cell(0, 1).text = "负责人"
    table.cell(1, 0).text = "发布"
    table.cell(1, 1).text = "产品"
    document.save(path)

    parsed = parse_file(path, filename="brief.docx")

    locators = [block["source_locator"] for block in parsed["blocks"]]
    assert locators[:2] == ["docx:paragraph=1", "docx:paragraph=3"]
    assert locators[-1] == "docx:table=1"
    assert parsed["blocks"][-1]["type"] == "table"


def test_xlsx_parser_emits_sheet_range_locator(tmp_path: Path):
    path = tmp_path / "metrics.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Revenue"
    sheet.append(["Month", "Amount"])
    sheet.append(["Jan", 10])
    sheet.append(["Feb", 20])
    workbook.save(path)

    parsed = parse_file(path, filename="metrics.xlsx")

    assert len(parsed["blocks"]) == 1
    assert parsed["blocks"][0]["source_locator"] == "xlsx:sheet=Revenue&range=A1:B3"
    assert parsed["blocks"][0]["row_count"] == 3


def test_text_parser_emits_stable_line_ranges(tmp_path: Path):
    path = tmp_path / "notes.txt"
    path.write_text("\n".join(f"line-{index}" for index in range(1, 206)), encoding="utf-8")

    parsed = parse_file(path, filename="notes.txt")

    assert [block["source_locator"] for block in parsed["blocks"]] == [
        "text:lines=1-200",
        "text:lines=201-205",
    ]


def test_image_and_empty_pdf_require_vision_instead_of_fake_success(tmp_path: Path):
    image_path = tmp_path / "scan.png"
    Image.new("RGB", (16, 16), "white").save(image_path)

    image_result = parse_file(image_path, mime_type="image/png")
    assert image_result["blocks"] == []
    assert image_result["metadata"]["processing_status"] == "vision_required"
    assert image_result["metadata"]["route_mode"] == "vision"

    pdf_path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    with pdf_path.open("wb") as output:
        writer.write(output)

    pdf_result = parse_file(pdf_path, mime_type="application/pdf")
    assert pdf_result["blocks"] == []
    assert pdf_result["metadata"]["processing_status"] == "vision_required"
