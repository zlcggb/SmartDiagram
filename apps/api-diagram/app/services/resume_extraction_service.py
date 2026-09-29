"""Resume text extraction: native text for PDF/DOCX/TXT, vision-LLM OCR for images and scans."""

from __future__ import annotations

import asyncio
import base64
import hashlib
import io
import re
import unicodedata
from collections import OrderedDict
from pathlib import Path
from typing import Any

from langchain_core.messages import HumanMessage

from app.core.llm import create_llm, extract_text_content
from app.services.recruit_ai_usage import record_usage, role_model, with_usage
from app.services.file_parser_service import DocumentParseError, parse_file

_MIN_PAGE_CHARS = 40
_MAX_OCR_IMAGES = 4
_MAX_IMAGE_SIDE = 2000
_OCR_CACHE: "OrderedDict[str, str]" = OrderedDict()
_OCR_CACHE_LIMIT = 128
_OCR_PROMPT = (
    "这是一份求职简历的页面图像。请完整、逐字转写页面上的全部文字，保持原有阅读顺序与分段：\n"
    "- 多栏排版按“先左栏自上而下、再右栏”的顺序输出；\n"
    "- 字段式信息把标签与对应的值写在同一行，每个字段单独一行，如 “姓名：张三”“籍贯：江苏扬州”；\n"
    "- 保留标题（如 教育背景 / 实习经历 / 项目经历 / 专业技能）独占一行，列表项用 “• ” 开头；\n"
    "- 教育经历每条写成一行：起止时间 学校 学历 专业；\n"
    "- 保留日期、数字、网址、邮箱原样；看不清的字用 □ 代替，不要猜测或补写；\n"
    "- 只输出转写文本，不要任何解释。"
)
_SPACED_TOKEN_RE = re.compile(r"(?<=\S) (?=\S)")
_LABEL_ONLY_LINE_RE = re.compile(r"^(?:[\u4e00-\u9fff]{2,4}\s*[:：]\s*){2,}$")


def _collapse_letter_spacing(text: str) -> str:
    """Undo per-glyph spacing (“教 育 背 景”); runs of 2+ spaces stay as word gaps."""

    lines: list[str] = []
    for line in text.splitlines():
        tokens = line.split()
        singles = sum(1 for token in tokens if len(token) == 1)
        if len(tokens) >= 4 and singles / len(tokens) >= 0.6:
            line = " ".join(_SPACED_TOKEN_RE.sub("", part) for part in re.split(r"\s{2,}", line.strip()))
        lines.append(line)
    return "\n".join(lines)


def _text_layer_garbled(text: str) -> bool:
    """Label grids split from their values (“姓名： 籍贯：”) or heavy glyph spacing mean the reading order is lost."""

    visible = [character for character in text if not character.isspace()]
    invalid = sum(
        character == "\ufffd" or unicodedata.category(character) in {"Cc", "Cf", "Cn", "Co", "Cs"}
        for character in visible
    )
    if any(character == "\ufffd" for character in visible) or (
        invalid >= 2 and invalid / max(len(visible), 1) >= 0.01
    ):
        return True

    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if not lines:
        return False
    label_only = sum(1 for line in lines if _LABEL_ONLY_LINE_RE.match(re.sub(r"(?<=\S) (?=\S)", "", line)))
    spaced = sum(1 for line in lines if len(line.split()) >= 4 and sum(len(t) == 1 for t in line.split()) / len(line.split()) >= 0.6)
    return label_only >= 1 or spaced / len(lines) >= 0.5


def _parse_pdfium_document(path: Path, *, mime_type: str, filename: str) -> dict[str, Any] | None:
    try:
        import pypdfium2 as pdfium
    except ImportError:
        return None

    document = None
    try:
        document = pdfium.PdfDocument(str(path))
        page_count = len(document)
        blocks: list[dict[str, Any]] = []
        for index in range(page_count):
            page = document[index]
            text_page = None
            try:
                text_page = page.get_textpage()
                text = (text_page.get_text_range() or "").strip()
            finally:
                if text_page is not None:
                    text_page.close()
                page.close()
            if text:
                blocks.append({
                    "type": "text",
                    "text": text,
                    "source_locator": f"pdf:page={index + 1}",
                    "page": index + 1,
                    "metadata": {"page": index + 1, "parser": "pdfium"},
                })
        if not blocks:
            return None
        return {
            "filename": filename or path.name,
            "mime_type": mime_type,
            "blocks": blocks,
            "metadata": {
                "suffix": ".pdf",
                "block_count": len(blocks),
                "size_bytes": path.stat().st_size,
                "parser": "pdfium",
                "route_mode": "full-context",
                "processing_status": "parsed",
                "page_count": page_count,
            },
        }
    except Exception:
        return None
    finally:
        if document is not None:
            try:
                document.close()
            except Exception:
                pass


def _render_pdf_pages(path: Path, page_numbers: list[int]) -> dict[int, bytes]:
    try:
        import pypdfium2 as pdfium
    except ImportError:
        return {}

    result: dict[int, bytes] = {}
    document = None
    try:
        document = pdfium.PdfDocument(str(path))
        for number in page_numbers:
            if number < 1 or number > len(document):
                continue
            page = document[number - 1]
            try:
                image = page.render(scale=2).to_pil()
                buffer = io.BytesIO()
                image.convert("RGB").save(buffer, format="JPEG", quality=88)
                result[number] = buffer.getvalue()
            except Exception:
                continue
            finally:
                try:
                    page.close()
                except Exception:
                    pass
    except Exception:
        return result
    finally:
        if document is not None:
            try:
                document.close()
            except Exception:
                pass
    return result


def _docx_textbox_blocks(path: Path, existing_text: str) -> list[dict[str, Any]]:
    """python-docx skips text boxes, which many Chinese resume templates rely on."""

    from docx import Document
    from docx.oxml.ns import qn

    document = Document(str(path))
    seen: set[str] = set()
    blocks: list[dict[str, Any]] = []
    for index, box in enumerate(document.element.body.iter(qn("w:txbxContent")), start=1):
        paragraphs = [
            "".join(node.text or "" for node in paragraph.iter(qn("w:t"))).strip()
            for paragraph in box.iter(qn("w:p"))
        ]
        text = "\n".join(p for p in paragraphs if p).strip()
        # VML fallback duplicates each DrawingML text box.
        if not text or text in seen or text in existing_text:
            continue
        seen.add(text)
        blocks.append({
            "type": "text",
            "text": text,
            "source_locator": f"docx:textbox={index}",
            "metadata": {"textbox_index": index},
        })
    return blocks


def _prepare_image(data: bytes) -> str | None:
    try:
        from PIL import Image

        image = Image.open(io.BytesIO(data))
        image.load()
        if image.width < 200 or image.height < 200:
            return None
        if image.mode not in ("RGB", "L"):
            image = image.convert("RGB")
        image.thumbnail((_MAX_IMAGE_SIDE, _MAX_IMAGE_SIDE))
        buffer = io.BytesIO()
        image.save(buffer, format="JPEG", quality=85)
        return "data:image/jpeg;base64," + base64.b64encode(buffer.getvalue()).decode("ascii")
    except Exception:
        return None


def _pdf_page_images(path: Path, page_numbers: list[int]) -> dict[int, bytes]:
    """Largest embedded raster per page — scanned PDFs store each page as one image."""

    from pypdf import PdfReader

    try:
        reader = PdfReader(str(path))
    except Exception:
        return {}
    result: dict[int, bytes] = {}
    for number in page_numbers:
        try:
            images = list(reader.pages[number - 1].images)
        except Exception:
            continue
        if images:
            largest = max(images, key=lambda item: len(item.data or b""))
            if largest.data:
                result[number] = largest.data
    return result


async def _ocr_image(data_url: str) -> str:
    llm = with_usage(create_llm(model=role_model("ocr"), temperature=0))
    response = await llm.ainvoke([
        HumanMessage(content=[
            {"type": "text", "text": _OCR_PROMPT},
            {"type": "image_url", "image_url": {"url": data_url}},
        ])
    ])
    text = extract_text_content(response.content).strip()
    record_usage("ocr", llm, response.usage_metadata, prompt_text=_OCR_PROMPT, output_text=text, images=1)
    return re.sub(r"^```\w*\s*|\s*```$", "", text).strip()


async def _ocr_many(images: list[bytes], *, allow_new: bool = True) -> list[str]:
    """OCR images, reusing cached results; without allow_new only the cache is consulted."""

    batch = images[:_MAX_OCR_IMAGES]
    digests = [hashlib.sha256(item).hexdigest() for item in batch]
    results: list[str] = [_OCR_CACHE.get(digest, "") for digest in digests]
    pending = [index for index, text in enumerate(results) if not text]
    if allow_new and pending:
        prepared = {index: _prepare_image(batch[index]) for index in pending}
        tasks = [_ocr_image(prepared[index]) for index in pending if prepared[index]]
        indexes = [index for index in pending if prepared[index]]
        outputs = await asyncio.gather(*tasks, return_exceptions=True)
        for index, output in zip(indexes, outputs):
            if isinstance(output, str) and output:
                results[index] = output
                _OCR_CACHE[digests[index]] = output
                while len(_OCR_CACHE) > _OCR_CACHE_LIMIT:
                    _OCR_CACHE.popitem(last=False)
    return results


def _quality(char_count: int, method: str) -> str:
    if char_count < 150:
        return "poor"
    if char_count < 600 or method == "ocr":
        return "fair"
    return "good"


def _resume_text_is_usable(text: str) -> bool:
    return len(text.strip()) >= _MIN_PAGE_CHARS and not _text_layer_garbled(text)


def resolve_resume_document(
    parsed_document: dict[str, Any] | None,
    manual_text: str | None = None,
) -> dict[str, Any] | None:
    if parsed_document is None:
        return None

    metadata = parsed_document.get("metadata") or {}
    extraction = metadata.get("extraction") or {}
    if (
        metadata.get("processing_status") not in {"review_required", "vision_required"}
        and extraction.get("verified") is not False
        and extraction.get("quality") != "poor"
    ):
        return parsed_document

    text = (manual_text or "").strip()
    if not _resume_text_is_usable(text):
        return None

    resolved_metadata = {
        **metadata,
        "block_count": 1,
        "processing_status": "parsed",
        "route_mode": "manual",
        "extraction": {
            "method": "manual",
            "char_count": len(text),
            "ocr_pages": [],
            "page_count": int(extraction.get("page_count") or metadata.get("page_count") or 0),
            "quality": "good" if len(text) >= 600 else "fair",
            "verified": True,
            "unresolved_pages": [],
            "warnings": ["简历正文由用户手动核对并提交。"],
        },
    }
    return {
        **parsed_document,
        "blocks": [{
            "type": "text",
            "text": text,
            "source_locator": "resume:manual",
            "metadata": {"manual": True},
        }],
        "metadata": resolved_metadata,
    }


async def extract_resume_document(
    path: str | Path,
    *,
    mime_type: str,
    filename: str,
    allow_ocr: bool = True,
) -> dict[str, Any]:
    """Parse a resume into locator-stable blocks, OCR-ing image-only content when allowed."""

    file_path = Path(path)
    suffix = (Path(filename).suffix if filename else file_path.suffix).lower()
    parsed = _parse_pdfium_document(file_path, mime_type=mime_type, filename=filename) if suffix == ".pdf" or mime_type == "application/pdf" else None
    if parsed is None:
        try:
            parsed = parse_file(file_path, mime_type=mime_type, filename=filename)
        except DocumentParseError:
            raise
        except Exception as exc:
            if suffix == ".pdf" or mime_type == "application/pdf":
                raise DocumentParseError("PDF 文件无法读取，请重新导出后重试。") from exc
            raise
    metadata = parsed["metadata"]
    suffix = str(metadata.get("suffix") or "")
    blocks = list(parsed["blocks"])
    warnings: list[str] = []
    ocr_pages: list[int] = []
    garbled_pages: list[int] = []
    unresolved_pages: list[int] = []
    document_unresolved = False

    if suffix == ".docx":
        try:
            blocks.extend(_docx_textbox_blocks(file_path, "\n".join(b.get("text", "") for b in blocks)))
        except Exception:
            document_unresolved = True
            warnings.append("Word 文本框内容读取失败，部分信息可能缺失。")

    if suffix == ".pdf":
        page_count = int(metadata.get("page_count") or 0)
        text_by_page = {int(b.get("page") or 0): len(b.get("text") or "") for b in blocks}
        sparse_pages = [n for n in range(1, page_count + 1) if text_by_page.get(n, 0) < _MIN_PAGE_CHARS]
        garbled_pages = sorted({
            int(b.get("page") or 0)
            for b in blocks
            if int(b.get("page") or 0) not in sparse_pages and _text_layer_garbled(str(b.get("text") or ""))
        })
        for block in blocks:
            block["text"] = _collapse_letter_spacing(str(block.get("text") or ""))
        review_pages = (sparse_pages + garbled_pages)[:_MAX_OCR_IMAGES]
        if review_pages:
            try:
                page_images = _pdf_page_images(file_path, [n for n in review_pages if n in sparse_pages])
            except Exception:
                page_images = {}
            missing = [n for n in review_pages if n not in page_images]
            if missing:
                try:
                    page_images.update(_render_pdf_pages(file_path, missing))
                except Exception:
                    pass
            if page_images:
                numbers = sorted(page_images)
                texts = await _ocr_many(
                    [page_images[n] for n in numbers],
                    allow_new=allow_ocr or any(n in garbled_pages for n in numbers),
                )
                for number, text in zip(numbers, texts):
                    if not _resume_text_is_usable(text):
                        continue
                    blocks = [b for b in blocks if int(b.get("page") or 0) != number]
                    blocks.append({
                        "type": "text",
                        "text": text,
                        "source_locator": f"pdf:page={number}:ocr",
                        "page": number,
                        "metadata": {"page": number, "ocr": True, "ai_review": number in garbled_pages},
                    })
                    ocr_pages.append(number)
                blocks.sort(key=lambda b: int(b.get("page") or 0))
            unresolved_pages = sorted((set(sparse_pages) | set(garbled_pages)) - set(ocr_pages))
            if unresolved_pages:
                blocks = [b for b in blocks if int(b.get("page") or 0) not in unresolved_pages]
            unresolved_sparse_pages = [n for n in sparse_pages if n in unresolved_pages]
            if unresolved_sparse_pages:
                warnings.append(f"PDF 有 {len(unresolved_sparse_pages)} 页未能识别出文字（可能是矢量图形或加密内容）。")
            if any(n in unresolved_pages for n in garbled_pages):
                warnings.append("PDF 文字层排版错乱且 AI 复核未成功，姓名等字段请人工核对。")

    image_unresolved = False
    if suffix in {".png", ".jpg", ".jpeg", ".webp"} or mime_type.startswith("image/"):
        texts = await _ocr_many([file_path.read_bytes()], allow_new=allow_ocr)
        if texts and _resume_text_is_usable(texts[0]):
            blocks = [{
                "type": "text",
                "text": texts[0],
                "source_locator": "image:ocr",
                "metadata": {"ocr": True},
            }]
            ocr_pages.append(1)
        else:
            image_unresolved = True
            warnings.append("图片中未识别出足够的文字，请确认图片清晰或改传 PDF / Word。")

    char_count = sum(len(b.get("text") or "") for b in blocks)
    native = any(not (b.get("metadata") or {}).get("ocr") for b in blocks)
    method = "mixed" if ocr_pages and native else "ocr" if ocr_pages else "text" if blocks else "none"
    quality = _quality(char_count, method)
    verified = (
        bool(blocks)
        and not unresolved_pages
        and not image_unresolved
        and not document_unresolved
        and method != "none"
        and quality != "poor"
    )
    if method == "none":
        warnings.append("未能从文件中提取到文字，请粘贴简历正文。")
    elif quality == "poor" and not unresolved_pages and not image_unresolved and not document_unresolved:
        warnings.append("提取到的文字过少，无法确认是完整简历；请人工核对。")
    parsed["blocks"] = blocks
    metadata.update({
        "block_count": len(blocks),
        "processing_status": "parsed" if verified else "review_required",
        "route_mode": "vision" if ocr_pages or unresolved_pages or image_unresolved else metadata.get("route_mode"),
        "extraction": {
            "method": method,
            "char_count": char_count,
            "ocr_pages": ocr_pages,
            "ai_review_pages": [n for n in ocr_pages if n in garbled_pages],
            "unresolved_pages": unresolved_pages or ([1] if image_unresolved else []),
            "page_count": int(metadata.get("page_count") or (1 if ocr_pages else 0)),
            "quality": quality if verified else "poor",
            "verified": verified,
            "warnings": warnings,
        },
    })
    return parsed
