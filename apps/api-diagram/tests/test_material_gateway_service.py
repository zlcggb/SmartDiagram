import io

import pytest
from pptx import Presentation
from starlette.datastructures import Headers, UploadFile

from app.services.material_gateway_service import (
    MAX_MATERIAL_BYTES,
    MaterialValidationError,
    sanitize_material_filename,
    stage_material_upload,
)


PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def test_default_material_upload_limit_is_50_mib():
    assert MAX_MATERIAL_BYTES == 50 * 1024 * 1024


def _pptx_bytes() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(0, 0, 100, 100)
    box.text = "季度经营复盘"
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _upload(filename: str, content: bytes, content_type: str) -> UploadFile:
    return UploadFile(
        io.BytesIO(content),
        filename=filename,
        headers=Headers({"content-type": content_type}),
    )


def test_sanitize_filename_keeps_unicode_basename_and_extension():
    safe = sanitize_material_filename("../../部门资料/ 季度\x00 经营复盘 .pptx")

    assert safe == "季度_经营复盘_.pptx"
    assert "/" not in safe
    assert "\x00" not in safe


@pytest.mark.asyncio
async def test_stage_upload_streams_and_validates_ooxml(tmp_path):
    content = _pptx_bytes()

    staged = await stage_material_upload(
        _upload("经营复盘.pptx", content, PPTX_MIME),
        temp_dir=tmp_path,
    )
    try:
        assert staged.safe_filename == "经营复盘.pptx"
        assert staged.mime_type == PPTX_MIME
        assert staged.size_bytes == len(content)
        assert len(staged.content_hash) == 64
        assert staged.route_mode == "full-context"
        assert staged.path.read_bytes() == content
    finally:
        staged.cleanup()

    assert not staged.path.exists()


@pytest.mark.asyncio
async def test_stage_upload_rejects_extension_magic_spoof(tmp_path):
    with pytest.raises(MaterialValidationError) as exc_info:
        await stage_material_upload(
            _upload("伪造.pptx", b"%PDF-1.7\n", PPTX_MIME),
            temp_dir=tmp_path,
        )

    assert exc_info.value.code == "file_signature_mismatch"
    assert exc_info.value.status_code == 415
    assert not list(tmp_path.iterdir())


@pytest.mark.asyncio
async def test_stage_upload_rejects_declared_mime_mismatch(tmp_path):
    with pytest.raises(MaterialValidationError) as exc_info:
        await stage_material_upload(
            _upload("复盘.pptx", _pptx_bytes(), "application/pdf"),
            temp_dir=tmp_path,
        )

    assert exc_info.value.code == "mime_type_mismatch"
    assert exc_info.value.status_code == 415


@pytest.mark.asyncio
async def test_stage_upload_rejects_empty_and_oversized_streams(tmp_path):
    with pytest.raises(MaterialValidationError) as empty_error:
        await stage_material_upload(
            _upload("empty.txt", b"", "text/plain"),
            temp_dir=tmp_path,
        )
    assert empty_error.value.code == "empty_file"

    with pytest.raises(MaterialValidationError) as large_error:
        await stage_material_upload(
            _upload("large.txt", b"a" * (MAX_MATERIAL_BYTES + 1), "text/plain"),
            temp_dir=tmp_path,
        )
    assert large_error.value.code == "file_too_large"
    assert large_error.value.status_code == 413
    assert not list(tmp_path.iterdir())


@pytest.mark.asyncio
async def test_image_upload_is_explicitly_routed_to_vision(tmp_path):
    png = b"\x89PNG\r\n\x1a\n" + b"\x00" * 32
    staged = await stage_material_upload(
        _upload("chart.png", png, "image/png"),
        temp_dir=tmp_path,
    )
    try:
        assert staged.route_mode == "vision"
        assert staged.processing_status == "vision_required"
    finally:
        staged.cleanup()
